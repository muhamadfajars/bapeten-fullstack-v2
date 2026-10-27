"""v072: source-backed, page-cohesive equipment + personnel identity rescue.

Bounded to the current document; does not restart OCR, alter preprocessing,
change address policy, or fabricate source values from the application form.
"""
from __future__ import annotations

from difflib import SequenceMatcher
from pathlib import Path
import re
import unicodedata


FACILITY_RE = re.compile(r'(?i)^nama\s+fasilitas(?:\s+kesehatan)?$')
BRAND_RE = re.compile(r'(?i)^alat\s+(\d+)\s*:\s*(?:merk|merek)\s+unit$')
UNIT_MODEL_RE = re.compile(r'(?i)^alat\s+(\d+)\s*:\s*(?:tipe|type|model)\s+unit$')
UNIT_SERIAL_RE = re.compile(r'(?i)^alat\s+(\d+)\s*:\s*(?:no\.?|nomor)\s+seri\s+unit$')
TUBE_SERIAL_RE = re.compile(r'(?i)^alat\s+(\d+)\s*:\s*(?:no\.?|nomor)\s+seri\s+(?:tabung|insersi\s+tabung)$')
GENERIC = {'', 'rumah sakit', 'rumah sakit jiwa', 'rs', 'klinik', 'none', '—', '-', 'null'}
DOCUMENT_TYPES = ('personel', 'surat izin', 'surat penunjukan', 'petugas tld')


def _text(value):
    return ' '.join(str(value or '').split())


def _simple(value):
    s = unicodedata.normalize('NFKC', _text(value)).casefold()
    return ' '.join(re.findall(r'[a-z0-9]+', s))


def _compact(value):
    return ''.join(re.findall(r'[a-z0-9]+', _simple(value)))


def _model(value):
    s = str(value or '')
    # Parenthetic product lines are not hardware model revisions.
    s = re.sub(r'\s*\(\s*[a-zA-Z][a-zA-Z ]+\s*\)\s*$', '', s)
    result = _compact(s)
    return result if len(result) >= 5 and re.search('[a-z]', result) and re.search('[0-9]', result) else ''


def _in_text(value, content):
    """Bounded ID match; punctuation in source model/serial is optional."""
    q = _compact(value)
    if len(q) < 5:
        return False
    pattern = r'(?<![A-Za-z0-9])' + r'[\s.\-/]*'.join(map(re.escape, q)) + r'(?![A-Za-z0-9])'
    return re.search(pattern, content, flags=re.I) is not None


def _pages(path: Path):
    """PDF text layer (or DOCX); no extra OCR pass and no cross-upload search."""
    try:
        if path.suffix.lower() == '.pdf':
            import fitz
            with fitz.open(str(path)) as pdf:
                return [(i + 1, pdf[i].get_text('text', sort=True)) for i in range(min(len(pdf), 60))]
        if path.suffix.lower() == '.docx':
            from docx import Document
            doc = Document(str(path))
            lines = [p.text for p in doc.paragraphs]
            for table in doc.tables[:35]:
                for row in table.rows[:120]:
                    lines.append(' | '.join(c.text for c in row.cells))
            return [(None, '\n'.join(lines))]
    except (ImportError, OSError, ValueError, RuntimeError):
        pass
    return []


def _clean_brand(value):
    value = _text(value).strip(' :;|\t')
    value = re.split(r'\s{3,}|\s*\|\s*|\s+(?:tipe\s+unit|no\.?\s+seri|model\s+tabung)\s*:', value, maxsplit=1, flags=re.I)[0]
    if not value or len(value) > 85 or len(value.split()) > 10 or re.match(r'(?i)^(?:tipe|type|nomor|no\.|model|merek|merk|serial|tahun|telp)', value):
        return ''
    return value


def _labeled_unit_brand(content):
    """Read Merek Unit from a real equipment data block; never manufacturer header."""
    lines = [_text(s) for s in content.splitlines() if _text(s)]
    pat = re.compile(r'(?i)^\s*(?:merek|merk)\s*(?:/\s*pabrikan\s*)?(?:unit|pesawat|generator)\s*[:\-]?\s*(.*)$')
    generic = re.compile(r'(?i)^\s*(?:merek|merk)\s*/\s*pabrikan\s*[:\-]?\s*(.+)$')
    outputs = []
    for i, line in enumerate(lines):
        m = pat.match(line) or generic.match(line)
        if not m:
            continue
        value = _clean_brand(m.group(1))
        if not value:
            for nxt in lines[i+1:i+3]:
                value = _clean_brand(nxt.lstrip(':').strip())
                if value:
                    break
        if value:
            outputs.append((value, line[:180]))
    return outputs


def _type_and_serial(rows, tool_id):
    found = {}
    for row in rows:
        if not isinstance(row, dict):
            continue
        field = _text(row.get('field'))
        for label, regex in (('model', UNIT_MODEL_RE), ('unit_serial', UNIT_SERIAL_RE), ('tube_serial', TUBE_SERIAL_RE)):
            m = regex.fullmatch(field)
            if m and m.group(1) == tool_id:
                v = row.get('observed') or row.get('expected')
                if v and _text(v).casefold() not in ('—','-','none'):
                    found[label] = v
    return found


def _anchor_pages(pages, rows, tool_id):
    """Anchor one device by unit model AND serial, on one or two nearby pages."""
    identifiers = _type_and_serial(rows, tool_id)
    if not identifiers.get('model') or not identifiers.get('unit_serial'):
        return []
    model = _model(identifiers['model'])
    serial = _compact(identifiers['unit_serial'])
    if not model or len(serial) < 6 or not re.search(r'\d', serial):
        return []
    scored=[]
    for i,(page,content) in enumerate(pages):
        if not isinstance(page,int):
            continue
        for other_page,other_content in [(-1,'')] + ([pages[i+1]] if i+1<len(pages) and pages[i+1][0]==page+1 else []):
            union = content + ('\n'+other_content if other_content else '')
            if not (_in_text(identifiers['unit_serial'],union) and _in_text(model,union)):
                continue
            # A two-page pair is permitted only when there is an actual
            # equipment section; otherwise these could be incidental IDs.
            if other_page!=-1 and not re.search(r'(?i)(?:data\s+pesawat|spesifikasi\s+alat|merk\s+unit|merek\s+unit|nomor\s+seri|serial\s+number)',union):
                continue
            score = 10 if other_page==-1 else 7
            if identifiers.get('tube_serial') and _in_text(identifiers['tube_serial'],union):
                score += 3
            if re.search(r'(?i)(?:data\s+pesawat|spesifikasi\s+alat|merk\s+unit|merek\s+unit)',union):
                score += 2
            scored.append((score,page))
    # Deduplicate page candidates, prefer earliest on equal evidence strength.
    seen={}
    for score,page in scored:
        seen[page] = max(score,seen.get(page,0))
    return sorted(((score,page) for page,score in seen.items()),key=lambda x:(-x[0],x[1]))


def refine_equipment(rows, pages):
    """Correct only Merek Unit with a brand in the same / neighboring section.

    Avoid using fallback hits from unrelated pages (e.g. vendor header p16).
    Preserve every other equipment field completely.
    """
    if not isinstance(rows, list):
        return rows
    for row in rows:
        if not isinstance(row, dict):
            continue
        match = BRAND_RE.fullmatch(_text(row.get('field')))
        if not match:
            continue
        ranked = _anchor_pages(pages, rows, match.group(1))
        if not ranked:
            continue
        page = ranked[0][1]
        local = [p for p, text in pages if p is not None and abs(p-page) <= 1 and _labeled_unit_brand(text)]
        # Always prefer the anchor page itself, then only directly adjacent pages.
        local.sort(key=lambda p: (abs(p-page), p))
        if not local:
            continue
        chosen = []
        for p in local:
            content = next(t for pg,t in pages if pg == p)
            chosen = [(p, brand, quote) for brand, quote in _labeled_unit_brand(content)]
            if chosen:
                break
        if len({_simple(b) for _,b,_ in chosen}) != 1:
            # Ambiguous multiple device brands in the same page => don't guess.
            continue
        p, brand, quote = chosen[0]
        if _simple(brand) in {'', 'merek unit', 'merk unit'}:
            continue
        old = _simple(row.get('observed'))
        # Do not reassign a definite correct match on a remote unrelated page
        # without authoritative same-device evidence (which we now have).
        row['observed'] = brand
        row['page'] = p
        row['quote'] = quote
        expected = _simple(row.get('expected'))
        actual = _simple(brand)
        if expected and expected == actual:
            row['status'], row['score'] = 'MATCH', 100.0
            row['note'] = 'Merek unit ditemukan dalam blok alat yang sama dengan model dan nomor seri.'
        elif expected and actual:
            row['status'] = 'MISMATCH'
            row['score'] = round(100*SequenceMatcher(None, expected, actual).ratio(),1)
            row['note'] = 'Merek unit berbeda pada blok alat yang terverifikasi oleh model dan nomor seri; periksa sumber.'
    return rows


def _canonical_facility(v):
    s = ' ' + unicodedata.normalize('NFKC', _text(v)).casefold() + ' '
    s = re.sub(r'\brumkit\b', ' rumah sakit tentara ', s)
    s = re.sub(r'\brsj\b', ' rumah sakit jiwa ', s)
    s = re.sub(r'\brs(?:u|ud)?\b', ' rumah sakit ', s)
    s = re.sub(r'\btk\.?\b', ' tingkat ', s)
    # Rank Roman numerals survive; internal military administrative code does not.
    s = re.sub(r'(?<=\s)\d{2}(?:[./]\d{2}){2}\b', ' ', s)
    words = re.findall(r'[a-z0-9]+', s)
    out=[]
    for w in words:
        if w == '4' and out and out[-1] == 'tingkat':w='iv'
        out.append(w)
    return out


def facility_equivalent(ref, candidate):
    a, b = _canonical_facility(ref), _canonical_facility(candidate)
    if not a or not b:
        return False
    # Strong explicit distinctive identity, not generic prefixes.
    generic = {'rumah','sakit','jiwa','tentara','tingkat','umum','daerah','klinik','rs','iv','i','ii','iii','v'}
    distinctive_a = [w for w in a if w not in generic]
    distinctive_b = [w for w in b if w not in generic]
    if not distinctive_a or not distinctive_b or distinctive_a != distinctive_b:
        return False
    if a == b:
        return True
    # Military hospital 'RS TK IV Pekanbaru' vs 'RUMKIT TK IV PEKANBARU'
    # may omit the adjective tentara; allow ONLY with same rank and locality.
    left = [w for w in a if w != 'tentara']
    right = [w for w in b if w != 'tentara']
    return (left == right and 'tingkat' in left and 'iv' in left and
            ('tentara' in a or 'tentara' in b))


def _facility_candidate(line, next_line=''):
    s = _text(line).strip(' \t|:;')
    if not s:
        return ''
    m = re.match(r'(?i)^(?:instansi|nama\s+(?:instansi|fasilitas|rumah\s+sakit)|tempat\s+praktik|customer(?:\s+name)?s?)\s*[:\-]?\s*(.+)$', s)
    if m:
        value=m.group(1).strip(' :|-')
        if value and len(value)<100:
            return value
    # Target practice / assignment identity, not the holder's private home address.
    practice = re.search(r'(?i)\b(?:di|untuk\s+melaksanakan\s+tugas\s+di)\s+((?:RSJ|RSUD|RSU|RS|RUMKIT|RUMAH\s+SAKIT)\.?\s+.+)',s)
    if practice:
        value=practice.group(1)
        value=re.split(r'(?i)\s+(?:(?:jl|jalan)\.?\s|sebagai\b|dengan\b|yang\b|untuk\b|no\.?\s*\d)',value,maxsplit=1)[0].strip(' ,.;:-')
        if 4<len(value)<100:return value
    return ''


def _source_names(pages):
    result=[]
    for page,content in pages:
        lines=[_text(x) for x in content.splitlines() if _text(x)]
        for i,line in enumerate(lines):
            candidate=_facility_candidate(line, lines[i+1] if i+1<len(lines) else '')
            if candidate:
                result.append((page,candidate,line[:260]))
    return result


def refine_personnel(rows, pages):
    """Rescue only weak/REVIEW facility entries from labeled employer/practice.

    Keep different explicit identities (MISMATCH) and all other fields intact.
    """
    if not isinstance(rows,list):
        return rows
    candidates = _source_names(pages)
    for row in rows:
        if not isinstance(row,dict) or not FACILITY_RE.fullmatch(_text(row.get('field'))):
            continue
        if str(row.get('status') or '').upper() == 'MISMATCH':
            continue
        observed = _simple(row.get('observed'))
        if observed and observed not in GENERIC and row.get('status')=='MATCH':
            continue
        # A nonempty contradictory name must not be replaced by an incidental
        # matching institution elsewhere in the same PDF.
        if observed not in GENERIC and not facility_equivalent(row.get('expected'),row.get('observed')):
            continue
        matches=[x for x in candidates if facility_equivalent(row.get('expected'),x[1])]
        if not matches:
            # An already extracted, explicit pair of equivalent military names
            # can be recognized without rewriting the source as a form value.
            if observed not in GENERIC and facility_equivalent(row.get('expected'),row.get('observed')):
                row['status']='MATCH';row['score']=100.0
                row['note']='Nama fasilitas setara setelah ekspansi singkatan terstruktur.'
            continue
        # Prefer Instansi/Nama Fasilitas line over casual free prose.
        matches.sort(key=lambda item:(0 if re.match(r'(?i)^(?:instansi|nama\s+)',item[2]) else 1,
                                      item[0] if isinstance(item[0],int) else 1))
        page,name,quote=matches[0]
        row['observed']=name
        row['status']='MATCH'
        row['score']=100.0
        if isinstance(page,int):row['page']=page
        row['quote']=quote
        row['note']='Nama fasilitas terbukti pada bagian instansi/tempat penugasan dalam dokumen.'
    return rows


def improve(rows, document_path):
    path = Path(document_path)
    if not isinstance(rows,list):
        return rows
    relevant=any(isinstance(r,dict) and (FACILITY_RE.fullmatch(_text(r.get('field'))) or BRAND_RE.fullmatch(_text(r.get('field')))) for r in rows)
    if not relevant:
        return rows
    pages = _pages(path)
    if not pages:
        return rows
    refine_equipment(rows,pages)
    # Use same facility policy in all categories; labels vs assignment context
    # are accepted only when strongly corroborated by the source.
    refine_personnel(rows,pages)
    return rows
